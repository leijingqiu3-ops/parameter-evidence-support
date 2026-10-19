# -*- coding: utf-8 -*-
"""Build a parameter-evidence DOCX from a JSON spec.

Usage:
    python build_evidence_docx.py spec.json output.docx
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None

ACCENT = "174A6E"
ACCENT_LIGHT = "EAF1F6"
ACCENT_PALE = "F4F8FA"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 103, 112)
RED = RGBColor(174, 58, 58)
GREEN = RGBColor(35, 108, 77)
AMBER = RGBColor(145, 94, 15)
FONT = "Microsoft YaHei"
LATIN = "Cambria"
DEFAULT_IMAGE_WIDTH_CM = 16.4
DEFAULT_IMAGE_MAX_HEIGHT_CM = 20.0
DEFAULT_IMAGE_PAGE_BREAK_HEIGHT_CM = 9.0

VERDICT_COLORS = {
    "keep": GREEN,
    "relabel": AMBER,
    "rederive": AMBER,
    "remove": RED,
    "missing": RED,
}


def set_run_font(run, size=None, bold=None, color=None, italic=None, latin=False):
    run.font.name = LATIN if latin else FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color) if isinstance(color, str) else color


def add_para(text, size=10.5, bold=False, color=TEXT, before=0, after=3,
             keep=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_label_value(label, value, value_color=TEXT, after=2, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    r1 = p.add_run(label)
    set_run_font(r1, size=9.8, bold=True, color=ACCENT)
    r2 = p.add_run(value)
    set_run_font(r2, size=9.8, color=value_color)
    return p


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FA8")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_table_borders(table, color="D8E0E6", size="4"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_formula(text, color=ACCENT):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.25)
    r = p.add_run(text)
    set_run_font(r, size=9.6, bold=False, color=color, latin=True)


def add_quote_block(label, quote):
    add_para(label, size=9.8, bold=True, color=ACCENT, before=2, after=2, keep=True)
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.4)
    cell.text = ""
    set_cell_shading(cell, ACCENT_PALE)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(quote)
    set_run_font(r, size=9.6, color=TEXT)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), ACCENT)
    borders.append(left)
    tcPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_entry(entry, number):
    doc.add_heading(f"{number}. {entry.get('heading', '未命名参数')}", level=2)
    for formula in entry.get("formulas", []):
        if isinstance(formula, (list, tuple)):
            add_formula(formula[0], RGBColor.from_string(formula[1].lstrip('#')) if len(formula) > 1 else ACCENT)
        else:
            add_formula(formula)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(entry.get("citation", ""))
    set_run_font(r, size=10.6, bold=True, color=TEXT)
    add_para(entry.get("journal", ""), size=9.5, color=MUTED, before=0, after=3, keep=True)
    add_label_value("论文标题：", entry.get("paper_title", ""), keep=True)
    add_label_value("作者：", entry.get("authors", ""), keep=True)
    add_label_value("链接类信息：", entry.get("links", ""), keep=True)
    add_quote_block("文献提及这个参数取值的正文段落：", entry.get("quote", ""))
    if entry.get("translation"):
        add_label_value("中文核对：", entry["translation"], after=3)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run("论文网页链接：")
    set_run_font(r1, size=10.2, bold=True, color=ACCENT)
    url = entry.get("url", "")
    if str(url).startswith("http"):
        add_hyperlink(p, url, url)
    else:
        r2 = p.add_run(url or "无外部实验论文")
        set_run_font(r2, size=9.8, color=MUTED)
    add_label_value("核验结论与建议改法：", entry.get("verdict", ""),
                    value_color=VERDICT_COLORS.get(entry.get("verdict_level", "keep"), TEXT), after=2)
    if entry.get("english"):
        add_label_value("可直接写入论文的英文表述：", entry["english"], after=2)
    local = entry.get("local", "")
    add_label_value("本地 PDF：", local or "【找不到全文】",
                    value_color=GREEN if local and "找不到" not in local else AMBER, after=4)
    add_snippet(entry)


def _as_image_items(entry):
    raw = entry.get("images")
    if raw is None:
        if entry.get("image"):
            raw = [{
                "path": entry.get("image"),
                "caption": entry.get("image_caption", ""),
                "width_cm": entry.get("image_width_cm"),
                "max_height_cm": entry.get("image_max_height_cm"),
                "page_break": entry.get("image_page_break"),
            }]
        else:
            raw = []
    elif isinstance(raw, (str, dict)):
        raw = [raw]
    return [{"path": item} if isinstance(item, str) else dict(item) for item in raw]


def _image_size_cm(path, width_cm, max_height_cm):
    try:
        width_cm = float(width_cm) if width_cm is not None else DEFAULT_IMAGE_WIDTH_CM
        max_height_cm = float(max_height_cm) if max_height_cm is not None else DEFAULT_IMAGE_MAX_HEIGHT_CM
    except (TypeError, ValueError):
        width_cm = DEFAULT_IMAGE_WIDTH_CM
        max_height_cm = DEFAULT_IMAGE_MAX_HEIGHT_CM
    if PILImage is None:
        return width_cm, None
    try:
        with PILImage.open(path) as image:
            px_width, px_height = image.size
    except Exception:
        return width_cm, None
    if px_width <= 0 or px_height <= 0:
        return width_cm, None
    height_cm = width_cm * px_height / px_width
    if max_height_cm > 0 and height_cm > max_height_cm:
        scale = max_height_cm / height_cm
        width_cm *= scale
        height_cm = max_height_cm
    return width_cm, height_cm


def _should_break_before_image(entry, item, height_cm):
    value = item.get("page_break", entry.get("image_page_break", "auto"))
    if isinstance(value, str) and value.lower() == "auto":
        threshold = entry.get(
            "image_page_break_height_cm",
            DEFAULT_IMAGE_PAGE_BREAK_HEIGHT_CM,
        )
        try:
            threshold = float(threshold)
        except (TypeError, ValueError):
            threshold = DEFAULT_IMAGE_PAGE_BREAK_HEIGHT_CM
        return height_cm is not None and height_cm > threshold
    return bool(value)


def _add_screenshot(entry, item, index, total):
    image = item.get("path") or item.get("image")
    prefix = f"原文截图 {index}：" if total > 1 else "原文截图："
    if not image:
        add_para(prefix + "【找不到截图】", size=9.2, color=AMBER, before=2, after=6)
        return
    path = Path(image)
    if not path.is_absolute():
        path = SPEC_DIR / path
    if not path.exists():
        add_para(prefix + f"【找不到截图】{image}", size=9.2, color=AMBER, before=2, after=6)
        return
    width_cm, height_cm = _image_size_cm(
        path,
        item.get("width_cm", entry.get("image_width_cm")),
        item.get("max_height_cm", entry.get("image_max_height_cm")),
    )
    if _should_break_before_image(entry, item, height_cm):
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    if height_cm is not None:
        run.add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
    else:
        run.add_picture(str(path), width=Cm(width_cm))
    caption = item.get("caption") or (entry.get("image_caption") if total == 1 else "")
    if caption:
        caption_para = add_para(caption, size=8.8, color=MUTED, before=0, after=8)
        caption_para.paragraph_format.keep_together = True


def add_snippet(entry):
    images = _as_image_items(entry)
    if not images:
        add_para("原文截图：【找不到截图】", size=9.2, color=AMBER, before=2, after=6)
        return
    for index, item in enumerate(images, 1):
        _add_screenshot(entry, item, index, len(images))


def add_page_number_footer(title):
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{title}  |  ")
    set_run_font(r, size=8.5, color=MUTED)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run = p.add_run(); run._r.append(f1); run._r.append(it); run._r.append(f2)
    set_run_font(run, size=8.5, color=MUTED)


def main():
    global doc, sec, SPEC_DIR
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    spec_path = Path(sys.argv[1]).resolve()
    out_path = Path(sys.argv[2]).resolve()
    SPEC_DIR = spec_path.parent
    spec = json.loads(spec_path.read_text(encoding="utf-8"))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(4)

    for name, size, color, before, after in [
        ("Title", 18, ACCENT, 0, 6),
        ("Subtitle", 10.5, MUTED, 0, 10),
        ("Heading 1", 16, ACCENT, 16, 8),
        ("Heading 2", 12.5, "233F54", 12, 5),
    ]:
        st = doc.styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color) if isinstance(color, str) else color
        st.font.bold = name != "Subtitle"
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = name in ("Heading 1", "Heading 2")

    title = spec.get("title", "参数文献依据核验")
    add_page_number_footer(title)

    p = doc.add_paragraph(style="Title")
    r = p.add_run(title)
    set_run_font(r, size=18, bold=True, color=ACCENT)
    p.paragraph_format.space_after = Pt(4)
    if spec.get("subtitle"):
        p = doc.add_paragraph(style="Subtitle")
        r = p.add_run(spec["subtitle"])
        set_run_font(r, size=10.5, color=MUTED)
    if spec.get("intro"):
        add_para(spec["intro"], after=6)

    if spec.get("grades"):
        add_para("证据分级", size=12, bold=True, color=ACCENT, before=5, after=4, keep=True)
        t = doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = [Cm(2.5), Cm(4.1), Cm(10.1)]
        for i, h in enumerate(["等级", "含义", "本文件中的例子"]):
            c = t.rows[0].cells[i]
            c.width = widths[i]
            c.text = ""
            p = c.paragraphs[0]
            r = p.add_run(h)
            set_run_font(r, size=9.6, bold=True, color=RGBColor(255, 255, 255))
            set_cell_shading(c, ACCENT)
            set_cell_margins(c)
        repeat_table_header(t.rows[0])
        for row in spec["grades"]:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].width = widths[i]
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                r = p.add_run(str(val))
                set_run_font(r, size=9.4, bold=(i == 0), color=TEXT if i else ACCENT)
                set_cell_margins(cells[i])
                if i == 0:
                    set_cell_shading(cells[i], ACCENT_LIGHT)
            prevent_row_split(t.rows[-1])
        set_table_borders(t)

    if spec.get("reading_notes"):
        add_para("阅读提示", size=11.5, bold=True, color=ACCENT, before=6, after=3, keep=True)
        for line in spec["reading_notes"]:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line)
            set_run_font(r, size=9.8, color=TEXT)

    if spec.get("revision_summary"):
        doc.add_page_break()
        doc.add_heading("修订清单", level=1)
        for block in spec["revision_summary"]:
            add_para(block.get("priority", "建议"), size=11.5, bold=True, color=ACCENT, before=6, after=3)
            for item in block.get("items", []):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(item)
                set_run_font(r, size=9.9, color=TEXT)

    number = 1
    for idx, section in enumerate(spec.get("sections", [])):
        if idx:
            doc.add_page_break()
        doc.add_heading(section.get("heading", ""), level=1)
        for entry in section.get("entries", []):
            add_entry(entry, number)
            number += 1

    if spec.get("not_found"):
        doc.add_page_break()
        doc.add_heading("找不到的文献与数值", level=1)
        add_para("以下条目未取得可核对的全文或截图；对应字段留空，不补写、不配替代引用。", after=6)
        t = doc.add_table(rows=1, cols=4)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = [Cm(4.0), Cm(4.6), Cm(5.1), Cm(2.4)]
        for i, h in enumerate(["文献", "待核对内容", "当前状态", "链接"]):
            c = t.rows[0].cells[i]
            c.width = widths[i]
            c.text = ""
            p = c.paragraphs[0]
            r = p.add_run(h)
            set_run_font(r, size=9.4, bold=True, color=RGBColor(255, 255, 255))
            set_cell_shading(c, ACCENT)
            set_cell_margins(c)
        repeat_table_header(t.rows[0])
        for item in spec["not_found"]:
            cells = t.add_row().cells
            for i, key in enumerate(["reference", "missing", "status"]):
                cells[i].width = widths[i]
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                r = p.add_run(item.get(key, ""))
                set_run_font(r, size=8.8, color=TEXT)
                set_cell_margins(cells[i])
            cells[3].width = widths[3]
            cells[3].text = ""
            p = cells[3].paragraphs[0]
            if str(item.get("url", "")).startswith("http"):
                add_hyperlink(p, item["url"], "DOI")
            set_cell_margins(cells[3])
            prevent_row_split(t.rows[-1])
        set_table_borders(t)

    if spec.get("files"):
        doc.add_page_break()
        add_para("已到位的核心 PDF", size=11.5, bold=True, color=ACCENT, before=0, after=3)
        for line in spec["files"]:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line)
            set_run_font(r, size=9.7, color=GREEN)

    if spec.get("advice"):
        add_para("使用建议", size=11.5, bold=True, color=ACCENT, before=8, after=3)
        for line in spec["advice"]:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line)
            set_run_font(r, size=10, color=TEXT)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(out_path)
    print("entries", number - 1)


if __name__ == "__main__":
    main()



